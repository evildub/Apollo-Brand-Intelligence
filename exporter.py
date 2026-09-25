import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict


# ── Required column layout ────────────────────────────────────────────────────
# A  = Title
# B  = URL
# C  = Thumbnail (image URL)
# D  = (blank spacer)
# E  = Item ID
# F  = (blank)
# G  = (blank)
# H  = Marketplace
# I  = (blank)
# J  = Seller Name
# K  = (blank)
# L  = (blank)
# M  = Brand
# N  = Price   ← extra, added at end of required cols

HEADERS = {
    "A": "Title",
    "B": "URL",
    "C": "Thumbnail",
    "D": "",
    "E": "Item ID",
    "F": "",
    "G": "",
    "H": "Marketplace",
    "I": "",
    "J": "Seller Name",
    "K": "",
    "L": "",
    "M": "Brand",
    "N": "Price",
    "O": "Item Location",
    "P": "Product Type",
    "Q": "Seller Origin (Registered)",
    "R": "Threat Assessment (3PL Hub / Origin)",
}

HEADER_FILL  = PatternFill("solid", fgColor="2B2D42")
HEADER_FONT  = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)

def normalize_price_string(price_str: str) -> str:
    """
    Sanitize and normalize price strings for Genesis Excel ingestion and UI consistency.
    Strips country prefixes (e.g. US $, USD $) while preserving international symbols (£, €, MXN, BRL, CAD, AUD).
    """
    if not price_str:
        return ""
    p = str(price_str).strip()
    if not p:
        return ""

    # 1. Strip promo / badge noise
    if any(w in p.lower() for w in ("free", "delivery", "shipping", "off", "%", "save")):
        return ""

    # 2. Reject watcher / quantity noise (e.g. "US 1", "US 3")
    if re.match(r'^(?:US|USD)?\s*\d{1,2}$', p, re.IGNORECASE):
        return ""

    # 3. Normalize "US $54.89" or "USD $54.89" -> "$54.89"
    p = re.sub(r'^(?:US|USD)\s*\$', '$', p, flags=re.IGNORECASE)
    p = re.sub(r'^(?:US|USD)\s+(?=\d)', '$', p, flags=re.IGNORECASE)

    # 4. If raw number (e.g. "54.89"), prepend standard "$"
    if re.match(r'^\d+(?:\.\d{2})?$', p):
        p = f"${p}"

    return p.strip()


def normalize_marketplace_code(mkt: str) -> str:
    """
    Ensure enterprise marketplace codes strictly adhere to client intake formatting and Genesis specifications.
    Strips all icons/emojis, whitespace, and normalizes canonical domain names (e.g. ebay.com, redbubble.com, printerval.com).
    """
    if not mkt:
        return MARKETPLACE

    raw_str = str(mkt).strip()
    # Remove leading/trailing emojis and non-alphanumeric symbols
    clean = re.sub(r'^[^\w\.\-]+|[^\w\.\-]+$', '', raw_str).strip()
    m = clean.lower()

    # 1. eBay & Regional Variations
    if any(k in m for k in ("cafr.ebay.ca", "cafr", "ca_fr", "canada (french)", "canada french", "ebay.ca (french)", "ebay.ca/cafr")):
        return "ebay.ca - cafr"
    if "ebay.co.uk" in m or ("uk" in m and "ebay" in m):
        return "ebay.co.uk"
    if "ebay.de" in m or ("germany" in m and "ebay" in m):
        return "ebay.de"
    if "ebay.ca" in m:
        return "ebay.ca"
    if "ebay.fr" in m:
        return "ebay.fr"
    if "ebay.it" in m:
        return "ebay.it"
    if "ebay.es" in m:
        return "ebay.es"
    if "ebay.com.au" in m or ("australia" in m and "ebay" in m):
        return "ebay.com.au"
    if "ebay.nl" in m:
        return "ebay.nl"
    if "ebay.pl" in m:
        return "ebay.pl"
    if "ebay.ch" in m:
        return "ebay.ch"
    if "ebay.at" in m:
        return "ebay.at"
    if "ebay.ie" in m:
        return "ebay.ie"
    if "ebay" in m:
        return "ebay.com"

    # 2. POD Platforms
    if "redbubble" in m:
        return "redbubble.com"
    if "printerval" in m:
        return "printerval.com"
    if "teepublic" in m:
        return "teepublic.com"
    if "spreadshirt" in m:
        return "spreadshirt.com"
    if "zazzle" in m:
        return "zazzle.com"
    if "etsy" in m:
        return "etsy.com"
    if "cafepress" in m:
        return "cafepress.com"
    if any(k in m for k in ("fineartamerica", "fine art america", "fineart", "pixels.com")):
        return "fineartamerica.com"
    if "threadless" in m:
        return "threadless.com"
    if any(k in m for k in ("teespring", "spring.com", "creator-spring")):
        return "teespring.com"

    # 3. Global Retail & Social Marketplaces
    if "tiktok" in m:
        return "shop.tiktok.com"
    if "aliexpress" in m or m == "ali":
        return "aliexpress.com"
    if "wish" in m:
        return "wish.com"
    if "temu" in m:
        return "temu.com"
    if "scribd" in m:
        return "scribd.com"

    # 4. ManoMano Locales
    if "manomano" in m:
        if "es" in m or "spain" in m:
            return "manomano.es"
        if "de" in m or "germany" in m:
            return "manomano.de"
        if "it" in m or "italy" in m:
            return "manomano.it"
        if "uk" in m or "co.uk" in m:
            return "manomano.co.uk"
        return "manomano.fr"

    # 5. Vinted Locales
    if "vinted" in m:
        if "fr" in m or "france" in m:
            return "vinted.fr"
        if "de" in m or "germany" in m:
            return "vinted.de"
        if "es" in m or "spain" in m:
            return "vinted.es"
        if "it" in m or "italy" in m:
            return "vinted.it"
        if "pl" in m or "poland" in m:
            return "vinted.pl"
        if "com" in m or "us" in m:
            return "vinted.com"
        return "vinted.co.uk"

    # 6. Mercado Libre Regional Domains
    if any(k in m for k in ("mercadolibre", "mercadolivre", "mercado", "meli")):
        if any(k in m for k in ("brazil", "brasil", "livre", "mlb", ".com.br")):
            return "mercadolivre.com.br"
        if any(k in m for k in ("argentina", "mla", ".com.ar")):
            return "mercadolibre.com.ar"
        if any(k in m for k in ("colombia", "mco", ".com.co")):
            return "mercadolibre.com.co"
        if any(k in m for k in ("chile", "mlc", ".cl")):
            return "mercadolibre.cl"
        if any(k in m for k in ("peru", "mpe", ".com.pe")):
            return "mercadolibre.com.pe"
        if any(k in m for k in ("uruguay", "mlu", ".com.uy")):
            return "mercadolibre.com.uy"
        return "listado.mercadolibre.com.mx"

    if "." in clean:
        return clean.lower()
    return f"{clean.lower()}.com" if clean else MARKETPLACE
HEADER_FILL  = PatternFill("solid", fgColor="2B2D42")
HEADER_FONT  = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
FONT_DATA_NORMAL = Font(name="Segoe UI", size=9)
FONT_DATA_BOLD   = Font(name="Segoe UI", size=9, bold=True)
FONT_DATA_LINK   = Font(color="0563C1", underline="single", name="Segoe UI", size=9)
FONT_REPEAT      = Font(name="Segoe UI", size=9, bold=True, color="B91C1C")
ALIGN_CENTER     = Alignment(horizontal="center", vertical="center")
ALIGN_DATA       = Alignment(vertical="center", wrap_text=False)
ALIGN_DATA_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=False)
ROW_FILL_A       = PatternFill("solid", fgColor="F8F8F2")
ROW_FILL_B       = PatternFill("solid", fgColor="EEEEEE")
BORDER_SIDE      = Side(style="thin", color="CCCCCC")
THIN_BORDER      = Border(bottom=BORDER_SIDE)

HEADER_COL_INDICES = {
    col: openpyxl.utils.column_index_from_string(col)
    for col in HEADERS
}

COL_WIDTHS = {
    "A": 55,   # Title
    "B": 50,   # URL
    "C": 45,   # Thumbnail
    "E": 16,   # Item ID
    "H": 14,   # Marketplace
    "J": 22,   # Seller
    "M": 18,   # Brand
    "N": 12,   # Price
    "O": 24,   # Item Location
    "P": 22,   # Product Type
    "Q": 22,   # Seller Origin
    "R": 38,   # Threat Assessment
}

MARKETPLACE = "ebay.com"


class ExcelExporter:
    def export(self, results: list[dict], filepath: str, progress_callback=None):
        """
        Export results to Excel.
        Tab 1: Master (All Results) containing 100% of harvested listings.
        Tabs 2+: Dedicated sheets per brand (grouped by item['brand']).
        """
        wb = openpyxl.Workbook()
        wb.properties.creator = "Jerry Seidenstucker"
        wb.properties.title = "Apollo Enforcement Harvester Report"
        wb.properties.description = "Generated by Apollo Brand Intelligence (Created by Jerry Seidenstucker)"
        wb.remove(wb.active)   # remove default sheet

        # 1. Tab 1: Master Sheet (All Listings)
        ws_all = wb.create_sheet(title="All Results")
        self._write_sheet(ws_all, results, "All Results", is_master=True, progress_callback=progress_callback)

        # 2. Tabs 2+: Dedicated Brand Tabs
        by_brand = defaultdict(list)
        for item in results:
            b = item.get("brand", "Unassigned")
            if not b or b == "Unknown":
                b = "Unassigned"
            by_brand[b].append(item)

        # Optimization: If there are > 25,000 rows and only 1 brand, writing a second identical sheet
        # doubles file size and doubles export time without adding any new data.
        # But if <= 25,000 or multiple brands exist, create dedicated brand sheets.
        if len(results) <= 25000 or len(by_brand) > 1:
            sorted_brands = sorted(by_brand.items(), key=lambda x: len(x[1]), reverse=True)[:50]
            for brand, items in sorted_brands:
                ws = wb.create_sheet(title=self._safe_sheet_name(brand))
                self._write_sheet(ws, items, brand, is_master=False)

        if progress_callback:
            progress_callback(len(results), len(results), "Saving Excel workbook to disk...")

        wb.save(filepath)
        return len(results)

    export_results = export

    # ── sheet writer ──────────────────────────────────────────────────────────
    def _write_sheet(self, ws, items: list[dict], brand: str, is_master: bool = False, progress_callback=None):
        ws.views.sheetView[0].showGridLines = True
        # header row
        for col_letter, header_text in HEADERS.items():
            col_idx = HEADER_COL_INDICES[col_letter]
            cell = ws.cell(row=1, column=col_idx, value=header_text)
            cell.font  = HEADER_FONT
            cell.fill  = HEADER_FILL
            cell.alignment = ALIGN_CENTER

        ws.row_dimensions[1].height = 20

        total_items = len(items)
        cb_step = max(250, total_items // 100) if total_items > 0 else 250

        # data rows (start at row 2)
        for idx, item in enumerate(items):
            row_num = idx + 2
            fill = ROW_FILL_A if row_num % 2 == 0 else ROW_FILL_B

            # Col A: Title
            cA = ws.cell(row=row_num, column=1, value=item.get("title", ""))
            cA.fill = fill
            cA.alignment = ALIGN_DATA
            cA.border = THIN_BORDER
            cA.font = FONT_DATA_NORMAL

            # Col B: URL
            url = item.get("url", "")
            cB = ws.cell(row=row_num, column=2, value=url)
            cB.fill = fill
            cB.alignment = ALIGN_DATA
            cB.border = THIN_BORDER
            if url:
                cB.hyperlink = url
                cB.font = FONT_DATA_LINK
            else:
                cB.font = FONT_DATA_NORMAL

            # Col C: Thumbnail
            img_url = item.get("image_url", "")
            cC = ws.cell(row=row_num, column=3, value=img_url)
            cC.fill = fill
            cC.alignment = ALIGN_DATA
            cC.border = THIN_BORDER
            if img_url:
                cC.hyperlink = img_url
                cC.font = FONT_DATA_LINK
            else:
                cC.font = FONT_DATA_NORMAL

            # Col D: Blank Spacer
            cD = ws.cell(row=row_num, column=4, value="")
            cD.fill = fill
            cD.alignment = ALIGN_DATA
            cD.border = THIN_BORDER
            cD.font = FONT_DATA_NORMAL

            # Col E: Item ID
            cE = ws.cell(row=row_num, column=5, value=item.get("item_id", ""))
            cE.fill = fill
            cE.alignment = ALIGN_DATA
            cE.border = THIN_BORDER
            cE.font = FONT_DATA_NORMAL

            # Col F: Blank Spacer
            cF = ws.cell(row=row_num, column=6, value="")
            cF.fill = fill
            cF.alignment = ALIGN_DATA
            cF.border = THIN_BORDER
            cF.font = FONT_DATA_NORMAL

            # Col G: Blank Spacer
            cG = ws.cell(row=row_num, column=7, value="")
            cG.fill = fill
            cG.alignment = ALIGN_DATA
            cG.border = THIN_BORDER
            cG.font = FONT_DATA_NORMAL

            # Col H: Marketplace
            cH = ws.cell(row=row_num, column=8, value=normalize_marketplace_code(item.get("marketplace", MARKETPLACE)))
            cH.fill = fill
            cH.alignment = ALIGN_DATA
            cH.border = THIN_BORDER
            cH.font = FONT_DATA_NORMAL

            # Col I: Blank Spacer
            cI = ws.cell(row=row_num, column=9, value="")
            cI.fill = fill
            cI.alignment = ALIGN_DATA
            cI.border = THIN_BORDER
            cI.font = FONT_DATA_NORMAL

            # Col J: Seller Name
            cJ = ws.cell(row=row_num, column=10, value=item.get("seller", ""))
            cJ.fill = fill
            cJ.alignment = ALIGN_DATA
            cJ.border = THIN_BORDER
            cJ.font = FONT_DATA_NORMAL

            # Col K: Blank Spacer
            cK = ws.cell(row=row_num, column=11, value="")
            cK.fill = fill
            cK.alignment = ALIGN_DATA
            cK.border = THIN_BORDER
            cK.font = FONT_DATA_NORMAL

            # Col L: Blank Spacer
            cL = ws.cell(row=row_num, column=12, value="")
            cL.fill = fill
            cL.alignment = ALIGN_DATA
            cL.border = THIN_BORDER
            cL.font = FONT_DATA_NORMAL

            # Col M: Brand
            brand_val = item.get("brand") or (brand if brand != "All Results" else "Unassigned")
            cM = ws.cell(row=row_num, column=13, value=brand_val)
            cM.fill = fill
            cM.alignment = ALIGN_DATA
            cM.border = THIN_BORDER
            cM.font = FONT_DATA_NORMAL

            # Col N: Price
            cN = ws.cell(row=row_num, column=14, value=normalize_price_string(item.get("price", "")))
            cN.fill = fill
            cN.alignment = ALIGN_DATA
            cN.border = THIN_BORDER
            cN.font = FONT_DATA_NORMAL

            # Col O: Item Location
            cO = ws.cell(row=row_num, column=15, value=item.get("location", ""))
            cO.fill = fill
            cO.alignment = ALIGN_DATA
            cO.border = THIN_BORDER
            cO.font = FONT_DATA_NORMAL

            # Col P: Product Type
            cP = ws.cell(row=row_num, column=16, value=item.get("product_type", ""))
            cP.fill = fill
            cP.alignment = ALIGN_DATA
            cP.border = THIN_BORDER
            cP.font = FONT_DATA_NORMAL

            # Col Q: Seller Origin
            cQ = ws.cell(row=row_num, column=17, value=item.get("seller_origin", item.get("country", "")))
            cQ.fill = fill
            cQ.alignment = ALIGN_DATA
            cQ.border = THIN_BORDER
            cQ.font = FONT_DATA_NORMAL

            # Col R: Threat Assessment
            cR = ws.cell(row=row_num, column=18, value=item.get("threat_badge", item.get("threat_intel", "")))
            cR.fill = fill
            cR.alignment = ALIGN_DATA
            cR.border = THIN_BORDER
            cR.font = FONT_DATA_NORMAL

            if progress_callback and (idx % cb_step == 0 or idx == total_items - 1):
                progress_callback(idx + 1, total_items, f"Writing rows ({idx + 1:,}/{total_items:,})...")

        # column widths
        for col_letter, width in COL_WIDTHS.items():
            ws.column_dimensions[col_letter].width = width

        # freeze header row
        ws.freeze_panes = "A2"

        # auto-filter
        ws.auto_filter.ref = f"A1:R1"

    def _safe_sheet_name(self, name: str) -> str:
        """Excel sheet names max 31 chars, no special chars."""
        invalid = r'\/?*[]:'
        for ch in invalid:
            name = name.replace(ch, "_")
        return name[:31]

    def export_a2c2_dossier(self, registry_dict: dict, file_path: str):
        """Export an executive A2C2 / Brand Protection Master Store & Recidivism Dossier."""
        wb = openpyxl.Workbook()
        ws_summary = wb.active
        ws_summary.title = "A2C2 Store Recidivism"
        ws_summary.views.sheetView[0].showGridLines = True

        headers = [
            "Seller / Store Name",
            "Offense Status",
            "Infringing Brands",
            "Product Types",
            "Total Infringing Listings",
            "Total Market Value ($ USD)",
            "Registered Origin",
            "Known 3PL Locations",
            "Cross-Border Threat Assessment",
            "First Detected",
            "Last Scanned",
            "Scan Count"
        ]

        ws_summary.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws_summary.cell(row=1, column=col_num)
            cell.fill = PatternFill("solid", fgColor="1E293B")
            cell.font = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
            cell.alignment = Alignment(horizontal="center" if col_num in (2, 5, 6, 12) else "left", vertical="center")

        row_idx = 2
        for seller, data in sorted(registry_dict.items(), key=lambda x: x[1].get("total_value", 0.0), reverse=True):
            scans = data.get("scan_count", 1)
            status = "🚨 REPEAT OFFENDER" if scans > 1 or data.get("total_listings", 0) > 10 else "⚠ FIRST STRIKE"
            brands = ", ".join(data.get("brands", []))
            pts = ", ".join(data.get("product_types", []))
            locs = ", ".join(data.get("locations", []))
            listings_count = data.get("total_listings", len(data.get("items", [])))
            total_val = f"${data.get('total_value', 0.0):,.2f}"
            origin = data.get("country") or data.get("seller_origin") or "Unresolved"
            threat = data.get("threat_badge") or data.get("threat_assessment") or ("🚨 OFFSHORE SMOKESCREEN" if "China" in origin else "Monitored")

            row_data = [
                seller,
                status,
                brands,
                pts,
                listings_count,
                total_val,
                origin,
                locs,
                threat,
                data.get("first_seen", ""),
                data.get("last_scanned", ""),
                scans
            ]
            ws_summary.append(row_data)

            fill = ROW_FILL_A if row_idx % 2 == 0 else ROW_FILL_B
            is_repeat = (status.startswith("🚨 REPEAT"))
            for col_num in range(1, len(headers) + 1):
                c = ws_summary.cell(row=row_idx, column=col_num)
                c.fill = fill
                c.font = FONT_REPEAT if (col_num == 2 and is_repeat) else FONT_DATA_NORMAL
                c.border = THIN_BORDER
                if col_num in (2, 5, 6, 12):
                    c.alignment = ALIGN_CENTER if col_num in (2, 12) else ALIGN_DATA_RIGHT
                else:
                    c.alignment = ALIGN_DATA
            row_idx += 1

        summary_widths = [24, 20, 26, 24, 22, 24, 20, 26, 32, 18, 18, 12]
        for idx, w in enumerate(summary_widths, 1):
            ws_summary.column_dimensions[get_column_letter(idx)].width = w
        ws_summary.freeze_panes = "A2"
        ws_summary.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

        all_items = []
        for seller, data in registry_dict.items():
            for it in data.get("items", []):
                all_items.append(it)

        if all_items:
            ws_items = wb.create_sheet(title="All Captured Listings")
            ws_items.views.sheetView[0].showGridLines = True
            item_headers = ["Seller", "Brand", "Product Type", "Item ID", "Title", "Price", "Seller Origin", "Location", "Threat Assessment", "Listing URL", "Scanned At"]
            ws_items.append(item_headers)
            for col_num in range(1, len(item_headers) + 1):
                cell = ws_items.cell(row=1, column=col_num)
                cell.fill = PatternFill("solid", fgColor="334155")
                cell.font = HEADER_FONT

            for i_idx, it in enumerate(all_items, start=2):
                orig_val = it.get("seller_origin") or it.get("country") or ""
                threat_val = it.get("threat_badge") or it.get("threat_intel") or ""
                ws_items.append([
                    it.get("seller", ""),
                    it.get("brand", ""),
                    it.get("product_type", ""),
                    it.get("item_id", ""),
                    it.get("title", ""),
                    it.get("price", ""),
                    orig_val,
                    it.get("location", ""),
                    threat_val,
                    it.get("url", ""),
                    it.get("scanned_at", "")
                ])
                fill = ROW_FILL_A if i_idx % 2 == 0 else ROW_FILL_B
                for col_num in range(1, len(item_headers) + 1):
                    c = ws_items.cell(row=i_idx, column=col_num)
                    c.fill = fill
                    c.font = FONT_DATA_NORMAL
                    c.border = THIN_BORDER
                    c.alignment = ALIGN_DATA

            item_widths = [18, 16, 20, 16, 45, 12, 18, 22, 30, 45, 18]
            for idx, w in enumerate(item_widths, 1):
                ws_items.column_dimensions[get_column_letter(idx)].width = w
            ws_items.freeze_panes = "A2"
            ws_items.auto_filter.ref = f"A1:{get_column_letter(len(item_headers))}1"

        wb.save(file_path)

    def export_multi_locale(self, results: list[dict], selected_locales: list[dict], file_path: str, progress_callback=None):
        """
        Export multi-locale projections formatted strictly to the 18-Column Genesis Upload Specification
        with Thumbnail in Col C and extended locale metadata starting in Column S.
        Tab 1: All Results (Master Multi-Locale Sheet)
        Tabs 2+: Dedicated sheets per brand
        """
        import re
        projected_rows = []
        for it in results:
            item_id = str(it.get("item_id", "")).strip()
            if not item_id:
                m = re.search(r"/itm/(\d+)", it.get("url", ""))
                if m:
                    item_id = m.group(1)
                else:
                    continue

            for loc in selected_locales:
                domain = loc.get("domain", "ebay.com")
                country = loc.get("name", "United States")
                flag = loc.get("flag", "🌍")
                region = loc.get("region", "Global")
                
                is_meli_item = "mercadolibre" in it.get("marketplace", "").lower() or "mercadolivre" in it.get("marketplace", "").lower() or item_id.startswith("ML")
                if is_meli_item:
                    if domain.startswith("http"):
                        locale_url = f"{domain}/articulo/{item_id}"
                    elif "listado" in domain or "lista" in domain:
                        locale_url = f"https://{domain}/{item_id}"
                    else:
                        locale_url = f"https://www.{domain}/articulo/{item_id}"
                elif domain == "cafr.ebay.ca":
                    locale_url = f"https://{domain}/itm/{item_id}"
                else:
                    locale_url = f"https://www.{domain}/itm/{item_id}"

                row_item = dict(it)
                row_item["item_id"] = item_id
                row_item["locale_url"] = locale_url
                row_item["locale_domain"] = normalize_marketplace_code(domain)
                row_item["locale_country"] = f"{flag} {country}"
                row_item["locale_region"] = region
                projected_rows.append(row_item)

        wb = openpyxl.Workbook()
        wb.properties.creator = "Jerry Seidenstucker"
        wb.properties.title = "Apollo Global Multi-Locale Enforcement Pack"
        wb.properties.description = "Generated by Apollo Brand Intelligence (Created by Jerry Seidenstucker)"
        wb.remove(wb.active)

        # 1. Master Tab (All Results)
        ws_all = wb.create_sheet(title="All Results")
        self._write_multi_locale_sheet(ws_all, projected_rows, progress_callback=progress_callback)

        # 2. Dedicated Brand Tabs
        by_brand = defaultdict(list)
        for r in projected_rows:
            b = r.get("brand", "Unassigned")
            if not b or b == "Unknown":
                b = "Unassigned"
            by_brand[b].append(r)

        # Optimization: If there are > 25,000 rows and only 1 brand, writing a second identical sheet
        # doubles file size and doubles export time without adding any new data.
        if len(projected_rows) <= 25000 or len(by_brand) > 1:
            sorted_brands = sorted(by_brand.items(), key=lambda x: len(x[1]), reverse=True)[:50]
            for brand, items in sorted_brands:
                ws_b = wb.create_sheet(title=self._safe_sheet_name(brand))
                self._write_multi_locale_sheet(ws_b, items)

        if progress_callback:
            progress_callback(len(projected_rows), len(projected_rows), "Saving Excel workbook to disk...")

        wb.save(file_path)
        return len(projected_rows)

    def _write_multi_locale_sheet(self, ws, items: list[dict], progress_callback=None):
        """Helper to write multi-locale rows into a worksheet with canonical Genesis styling."""
        ws.views.sheetView[0].showGridLines = True

        headers_map = {
            "A": "Title",
            "B": "URL",
            "C": "Thumbnail",
            "D": "",
            "E": "Item ID",
            "F": "",
            "G": "",
            "H": "Marketplace",
            "I": "",
            "J": "Seller Name",
            "K": "",
            "L": "",
            "M": "Brand",
            "N": "Price",
            "O": "Item Location",
            "P": "Product Type",
            "Q": "Seller Origin (Registered)",
            "R": "Threat Assessment (3PL Hub / Origin)",
            "S": "Locale Country",
            "T": "Locale Domain",
            "U": "Region"
        }

        # Write Header Row
        for col_idx, (col_letter, header_text) in enumerate(headers_map.items(), start=1):
            cell = ws.cell(row=1, column=col_idx, value=header_text)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = ALIGN_CENTER

        ws.row_dimensions[1].height = 22

        total_items = len(items)
        cb_step = max(250, total_items // 100) if total_items > 0 else 250

        for idx, it in enumerate(items):
            row_idx = idx + 2
            fill = ROW_FILL_A if row_idx % 2 == 0 else ROW_FILL_B

            # Col A: Title
            cA = ws.cell(row=row_idx, column=1, value=it.get("title", ""))
            cA.fill = fill
            cA.alignment = ALIGN_DATA
            cA.border = THIN_BORDER
            cA.font = FONT_DATA_NORMAL

            # Col B: URL
            loc_url = it.get("locale_url") or it.get("url", "")
            cB = ws.cell(row=row_idx, column=2, value=loc_url)
            cB.fill = fill
            cB.alignment = ALIGN_DATA
            cB.border = THIN_BORDER
            if loc_url:
                cB.hyperlink = loc_url
                cB.font = FONT_DATA_LINK
            else:
                cB.font = FONT_DATA_NORMAL

            # Col C: Thumbnail
            img_url = it.get("image_url", "")
            cC = ws.cell(row=row_idx, column=3, value=img_url)
            cC.fill = fill
            cC.alignment = ALIGN_DATA
            cC.border = THIN_BORDER
            if img_url:
                cC.hyperlink = img_url
                cC.font = FONT_DATA_LINK
            else:
                cC.font = FONT_DATA_NORMAL

            # Col D: Blank Spacer
            cD = ws.cell(row=row_idx, column=4, value="")
            cD.fill = fill
            cD.alignment = ALIGN_DATA
            cD.border = THIN_BORDER
            cD.font = FONT_DATA_NORMAL

            # Col E: Item ID
            cE = ws.cell(row=row_idx, column=5, value=it.get("item_id", ""))
            cE.fill = fill
            cE.alignment = ALIGN_DATA
            cE.border = THIN_BORDER
            cE.font = FONT_DATA_NORMAL

            # Col F: Blank Spacer
            cF = ws.cell(row=row_idx, column=6, value="")
            cF.fill = fill
            cF.alignment = ALIGN_DATA
            cF.border = THIN_BORDER
            cF.font = FONT_DATA_NORMAL

            # Col G: Blank Spacer
            cG = ws.cell(row=row_idx, column=7, value="")
            cG.fill = fill
            cG.alignment = ALIGN_DATA
            cG.border = THIN_BORDER
            cG.font = FONT_DATA_NORMAL

            # Col H: Marketplace
            cH = ws.cell(row=row_idx, column=8, value=it.get("locale_domain") or normalize_marketplace_code(it.get("marketplace", MARKETPLACE)))
            cH.fill = fill
            cH.alignment = ALIGN_DATA
            cH.border = THIN_BORDER
            cH.font = FONT_DATA_NORMAL

            # Col I: Blank Spacer
            cI = ws.cell(row=row_idx, column=9, value="")
            cI.fill = fill
            cI.alignment = ALIGN_DATA
            cI.border = THIN_BORDER
            cI.font = FONT_DATA_NORMAL

            # Col J: Seller
            cJ = ws.cell(row=row_idx, column=10, value=it.get("seller", ""))
            cJ.fill = fill
            cJ.alignment = ALIGN_DATA
            cJ.border = THIN_BORDER
            cJ.font = FONT_DATA_NORMAL

            # Col K: Blank Spacer
            cK = ws.cell(row=row_idx, column=11, value="")
            cK.fill = fill
            cK.alignment = ALIGN_DATA
            cK.border = THIN_BORDER
            cK.font = FONT_DATA_NORMAL

            # Col L: Blank Spacer
            cL = ws.cell(row=row_idx, column=12, value="")
            cL.fill = fill
            cL.alignment = ALIGN_DATA
            cL.border = THIN_BORDER
            cL.font = FONT_DATA_NORMAL

            # Col M: Brand
            cM = ws.cell(row=row_idx, column=13, value=it.get("brand", ""))
            cM.fill = fill
            cM.alignment = ALIGN_DATA
            cM.border = THIN_BORDER
            cM.font = FONT_DATA_NORMAL

            # Col N: Price
            cN = ws.cell(row=row_idx, column=14, value=it.get("price", ""))
            cN.fill = fill
            cN.alignment = ALIGN_DATA
            cN.border = THIN_BORDER
            cN.font = FONT_DATA_NORMAL

            # Col O: Location
            cO = ws.cell(row=row_idx, column=15, value=it.get("location", ""))
            cO.fill = fill
            cO.alignment = ALIGN_DATA
            cO.border = THIN_BORDER
            cO.font = FONT_DATA_NORMAL

            # Col P: Product Type
            cP = ws.cell(row=row_idx, column=16, value=it.get("product_type", ""))
            cP.fill = fill
            cP.alignment = ALIGN_DATA
            cP.border = THIN_BORDER
            cP.font = FONT_DATA_NORMAL

            # Col Q: Seller Origin
            cQ = ws.cell(row=row_idx, column=17, value=it.get("seller_origin", it.get("country", "")))
            cQ.fill = fill
            cQ.alignment = ALIGN_DATA
            cQ.border = THIN_BORDER
            cQ.font = FONT_DATA_NORMAL

            # Col R: Threat Assessment
            cR = ws.cell(row=row_idx, column=18, value=it.get("threat_badge", it.get("threat_intel", "")))
            cR.fill = fill
            cR.alignment = ALIGN_DATA
            cR.border = THIN_BORDER
            cR.font = FONT_DATA_NORMAL

            # Col S: Locale Country
            cS = ws.cell(row=row_idx, column=19, value=it.get("locale_country", ""))
            cS.fill = fill
            cS.alignment = ALIGN_DATA
            cS.border = THIN_BORDER
            cS.font = FONT_DATA_NORMAL

            # Col T: Locale Domain
            cT = ws.cell(row=row_idx, column=20, value=it.get("locale_domain", ""))
            cT.fill = fill
            cT.alignment = ALIGN_DATA
            cT.border = THIN_BORDER
            cT.font = FONT_DATA_NORMAL

            # Col U: Locale Region
            cU = ws.cell(row=row_idx, column=21, value=it.get("locale_region", ""))
            cU.fill = fill
            cU.alignment = ALIGN_DATA
            cU.border = THIN_BORDER
            cU.font = FONT_DATA_NORMAL

            if progress_callback and (idx % cb_step == 0 or idx == total_items - 1):
                progress_callback(idx + 1, total_items, f"Writing rows ({idx + 1:,}/{total_items:,})...")

        col_widths = {
            "A": 55, "B": 50, "C": 45, "D": 4, "E": 16, "F": 4, "G": 4,
            "H": 18, "I": 4, "J": 22, "K": 4, "L": 4, "M": 18, "N": 12,
            "O": 24, "P": 20, "Q": 20, "R": 30, "S": 22, "T": 18, "U": 16
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:U1"

