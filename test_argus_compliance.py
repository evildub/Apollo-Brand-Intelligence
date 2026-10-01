"""
test_argus_compliance.py - Comprehensive Unit Tests for Argus Compliance Sentinel
Verifies marketplace routing, status classification heuristics, concurrency controls,
file ingestion, Excel export formatting, and UI lifecycle.
"""

import os
import sys
import unittest
import tempfile
import csv
from unittest.mock import patch, MagicMock

import tkinter as tk

from argus_engine import (
    ArgusEngine,
    ComplianceResult,
    STATUS_REMOVED,
    STATUS_ENDED,
    STATUS_ACTIVE,
    STATUS_BLOCKED,
    STATUS_ERROR,
    STATUS_BADGES
)


class TestArgusEngine(unittest.TestCase):
    def setUp(self):
        self.engine = ArgusEngine(max_workers=4, timeout=5)
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        if os.path.exists(self.temp_dir):
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_marketplace_identification(self):
        """Verify standard domain parsing maps accurately to supported marketplaces."""
        cases = [
            ("https://www.ebay.com/itm/1234567890", "eBay"),
            ("https://www.ebay.co.uk/itm/987654321", "eBay"),
            ("https://articulo.mercadolibre.com.mx/MLM-123456", "Mercado Libre"),
            ("https://produto.mercadolivre.com.br/MLB-987654", "Mercado Libre"),
            ("https://www.amazon.com/dp/B08N5WRWNW", "Amazon"),
            ("https://www.redbubble.com/i/t-shirt/Cool-Design/12345.1YY9G", "Redbubble"),
            ("https://www.teepublic.com/t-shirt/998877-cool-shirt", "TeePublic"),
            ("https://www.vinted.fr/items/12345678-vintage-hoodie", "Vinted"),
            ("https://www.etsy.com/listing/11223344/handcrafted-ring", "Etsy"),
            ("https://www.aliexpress.com/item/100500123456.html", "AliExpress"),
            ("https://printblur.com/product/12345", "Printblur"),
            ("https://printerval.com/custom-hoodie-p12345", "Printerval"),
            ("https://customstore.myshopify.com/products/item-1", "Shopify"),
            ("https://unknown-shop.net/item/123", "unknown-shop.net")
        ]
        for url, expected in cases:
            mkt = self.engine.identify_marketplace(url)
            self.assertEqual(mkt, expected, f"Failed for URL: {url}")

    @patch("argus_engine.curl_requests.get")
    def test_probe_404_removed(self, mock_get):
        """Verify HTTP 404 triggers REMOVED / 404 status."""
        mock_resp = MagicMock()
        mock_resp.status_code = 404
        mock_resp.text = "<html><body>404 Not Found</body></html>"
        mock_get.return_value = mock_resp

        res = self.engine.probe_url("https://www.ebay.com/itm/12345")
        self.assertEqual(res.status, STATUS_REMOVED)
        self.assertEqual(res.status_code, 404)
        self.assertIn("Removed", res.details)

    @patch("argus_engine.curl_requests.get")
    def test_ebay_ended_heuristic(self, mock_get):
        """Verify eBay 'listing was ended by the seller' heuristic triggers ENDED status."""
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.text = "<html><body>This listing was ended by the seller because the item was lost or broken.</body></html>"
        mock_get.return_value = mock_resp

        res = self.engine.probe_url("https://www.ebay.com/itm/12345")
        self.assertEqual(res.status, STATUS_ENDED)
        self.assertEqual(res.marketplace, "eBay")
        self.assertIn("ended by seller", res.details.lower())

    @patch("argus_engine.curl_requests.get")
    def test_mercadolibre_finalizada_heuristic(self, mock_get):
        """Verify Mercado Libre 'Publicación finalizada' heuristic triggers ENDED status."""
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.text = "<html><body><h1>Publicación finalizada</h1><p>El vendedor finalizó esta publicación.</p></body></html>"
        mock_get.return_value = mock_resp

        res = self.engine.probe_url("https://articulo.mercadolibre.com.mx/MLM-12345")
        self.assertEqual(res.status, STATUS_ENDED)
        self.assertEqual(res.marketplace, "Mercado Libre")

    @patch("argus_engine.curl_requests.get")
    def test_redbubble_removed_heuristic(self, mock_get):
        """Verify Redbubble 'This work has been removed' triggers REMOVED status."""
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.text = "<html><body>Whoops! We couldn't find that page. This work has been removed.</body></html>"
        mock_get.return_value = mock_resp

        res = self.engine.probe_url("https://www.redbubble.com/i/t-shirt/test/123")
        self.assertEqual(res.status, STATUS_REMOVED)
        self.assertEqual(res.marketplace, "Redbubble")

    @patch("argus_engine.curl_requests.get")
    def test_active_buy_button_heuristic(self, mock_get):
        """Verify active buy button triggers ACTIVE status."""
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.text = "<html><body><button id='binBtn_btn'>Buy It Now</button><button>Add to cart</button><div>In Stock</div></body></html>" * 200
        mock_get.return_value = mock_resp

        res = self.engine.probe_url("https://www.ebay.com/itm/99999")
        self.assertEqual(res.status, STATUS_ACTIVE)
        self.assertEqual(res.status_code, 200)

    def test_parse_csv_file(self):
        """Verify CSV import correctly auto-identifies URL column and preserves fields."""
        csv_file = os.path.join(self.temp_dir, "test_input.csv")
        with open(csv_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Brand", "Item ID", "Listing URL", "Seller", "Price"])
            writer.writerow(["Toyota", "12345", "https://www.ebay.com/itm/12345", "speedy_seller", "$25.00"])
            writer.writerow(["Ford", "67890", "https://www.amazon.com/dp/B0001", "auto_parts_direct", "$40.00"])

        rows, url_col, all_cols = ArgusEngine.parse_import_file(csv_file)
        self.assertEqual(len(rows), 2)
        self.assertEqual(url_col, "Listing URL")
        self.assertEqual(rows[0]["Brand"], "Toyota")
        self.assertEqual(rows[1]["Seller"], "auto_parts_direct")

    def test_export_compliance_report(self):
        """Verify Excel compliance report generation with openpyxl formatting."""
        results = [
            ComplianceResult("https://www.ebay.com/itm/123", "eBay", STATUS_REMOVED, 404, "Page Not Found", 120.5, {"Brand": "Toyota", "Seller": "seller_a"}),
            ComplianceResult("https://www.ebay.com/itm/456", "eBay", STATUS_ENDED, 200, "Ended by seller", 95.0, {"Brand": "Toyota", "Seller": "seller_b"}),
            ComplianceResult("https://www.amazon.com/dp/789", "Amazon", STATUS_ACTIVE, 200, "In Stock", 150.0, {"Brand": "Ford", "Seller": "seller_c"})
        ]

        out_path = os.path.join(self.temp_dir, "argus_report.xlsx")
        ArgusEngine.export_compliance_report(out_path, results, original_cols=["Brand", "Seller"])

        self.assertTrue(os.path.exists(out_path))
        self.assertGreater(os.path.getsize(out_path), 1000)

        # Re-read with openpyxl to verify sheet structure
        import openpyxl
        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        self.assertEqual(ws.title, "Argus Compliance Audit")
        self.assertEqual(ws.cell(row=1, column=1).value, "Compliance Status")
        self.assertEqual(ws.cell(row=2, column=1).value, STATUS_BADGES[STATUS_REMOVED])
        self.assertEqual(ws.cell(row=3, column=1).value, STATUS_BADGES[STATUS_ENDED])
        self.assertEqual(ws.cell(row=4, column=1).value, STATUS_BADGES[STATUS_ACTIVE])
        wb.close()

    def test_concurrency_and_abort(self):
        """Verify batch auditing responds properly to abort event."""
        items = [{"url": f"https://www.ebay.com/itm/{i}"} for i in range(20)]
        self.engine.abort()  # Pre-abort
        results = self.engine.audit_batch(items)
        self.assertEqual(len(results), 0)


class TestArgusModalUI(unittest.TestCase):
    def setUp(self):
        self.root = tk.Tk()
        self.root.withdraw()

    def tearDown(self):
        try:
            self.root.destroy()
        except Exception:
            pass

    def test_modal_init_and_elements(self):
        """Verify ArgusComplianceModal initializes without error and renders all controls."""
        from argus_compliance_modal import ArgusComplianceModal
        modal = ArgusComplianceModal(master=self.root)
        self.assertTrue(modal.winfo_exists())
        self.assertIsNotNone(modal.tree)
        self.assertIsNotNone(modal.btn_start)
        self.assertIsNotNone(modal.btn_export)
        self.assertIsNotNone(modal.prog_bar)
        modal.destroy()


if __name__ == "__main__":
    unittest.main()
