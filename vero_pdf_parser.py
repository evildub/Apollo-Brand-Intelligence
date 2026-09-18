"""
vero_pdf_parser.py
Specialized VeRO & eBay Seller Disclosure Intelligence Harvester & Parser for Apollo.

Features:
- Extracts raw seller tracking text lines from VeRO / eBay PDF disclosure documents or direct text pastes.
- Intelligently decomposes disclosure text lines into structured components:
  * Seller Handle
  * Contact / Legal Name (Chinese Pinyin, Western names, etc.)
  * Street Address / Building / Unit
  * City / Municipal District
  * Postal / Zip Code
  * Country Code (ISO-2) / Region
- Ingests and updates records directly into Apollo's permanent Enforcement Registry (`data_store.py`).
- Generates Genesis-compliant Excel (.xlsx) workbooks and CSV spreadsheets ready for CSM and client intake.
"""

import os
import re
import csv
import logging
from datetime import datetime
from typing import List, Dict, Optional, Any

try:
    from pypdf import PdfReader
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger("Apollo.VeroParser")

ADDR_SUFFIXES = (
    "qu", "lu", "dao", "jie", "ceng", "fang", "cheng", "road", "street", "st", "rd",
    "ave", "blvd", "way", "lane", "drive", "dr", "court", "ct", "place", "pl",
    "strasse", "straße", "str", "weg", "gasse", "rue", "calle", "via", "corso", "alley"
)

ADDR_KEYWORDS = {
    "room", "suite", "apt", "unit", "building", "district", "qu", "lu", "no", "no.",
    "haus", "bldg", "fl", "floor", "ste", "block", "lot", "dangkouhao"
}


def parse_vero_line(line: str) -> Optional[Dict[str, str]]:
    """
    Parse a single VeRO / eBay seller disclosure line into structured components.
    Handles US, European, and Chinese Pinyin multi-segment addresses without column shifts.
    Example: 'trdracing / xu jie xu yu xiu qu yong fu lu 35 2, guang zhou, 510000, CN'
    """
    if not line or not line.strip():
        return None

    raw = line.strip()

    # 1. Split Seller Handle from Disclosure Body
    parts = re.split(r'\s+/\s+|\s*\|\s*|\t+|\s*:\s*(?=[a-zA-Z0-9])', raw, maxsplit=1)
    if len(parts) == 2:
        handle = parts[0].strip()
        disclosure = parts[1].strip()
    else:
        tokens = raw.split()
        if len(tokens) >= 2 and ('/' in tokens[0] or '_' in tokens[0] or tokens[0].isalnum()):
            handle = tokens[0].strip('/')
            disclosure = " ".join(tokens[1:]).strip()
        else:
            handle = ""
            disclosure = raw

    # Clean handle
    handle = re.sub(r'^(?:seller|user|handle|id)[:\s]*', '', handle, flags=re.IGNORECASE).strip()

    # 2. Split comma segments
    comma_segments = [s.strip() for s in disclosure.split(',') if s.strip()]

    country = ""
    postal_code = ""
    city = ""
    contact_name = ""
    street_address = ""

    # Common country identifiers
    country_patterns = r'^(?:CN|US|USA|GB|UK|DE|FR|CA|AU|IT|ES|JP|HK|TW|KR|China|United States|United Kingdom|Germany|France|Canada|Australia|Hong Kong|Italy|Spain|Japan|Taiwan|South Korea)$'

    # Work backwards from the rightmost segments
    rem = list(comma_segments)

    # Detect country in last segment
    if rem:
        last = rem[-1]
        if len(last) <= 3 or re.match(country_patterns, last, re.IGNORECASE) or (len(last) <= 20 and not any(c.isdigit() for c in last)):
            country = rem.pop()

    # Detect Postal Code / State / City from remaining trailing segments
    if rem:
        seg = rem[-1]
        # Check if segment has "State Zip" (e.g., "TX 78701", "CA 90210-1234", "ON M5V 2T6")
        state_zip_m = re.search(r'\b([A-Z]{2})\s+([A-Z0-9\s-]{3,10})$', seg, re.IGNORECASE)
        # Or purely a zip code (e.g. "510000", "78701", "90210-1234", "SW1A 1AA")
        pure_zip_m = re.match(r'^(?:\d{4,8}(?:-\d{4})?|[A-Z0-9]{2,4}\s?[A-Z0-9]{2,4})$', seg, re.IGNORECASE)

        if state_zip_m:
            postal_code = state_zip_m.group(2).strip()
            leftover = seg[:state_zip_m.start()].strip()
            rem.pop()
            if leftover:
                rem.append(leftover)
        elif pure_zip_m:
            postal_code = rem.pop()

    # Next trailing segment is usually City / District
    if rem and len(rem) >= 2:
        city = rem.pop()
    elif rem and len(rem) == 1 and not city:
        # Check if single remaining token has City at the end
        pass

    # If postal code was embedded in city (e.g. "guang zhou 510000" or "Austin TX 78701")
    if city and not postal_code:
        zip_in_city = re.search(r'\b(\d{4,8}(?:-\d{4})?|[A-Z]\d[A-Z]\s?\d[A-Z]\d)\b', city)
        if zip_in_city:
            postal_code = zip_in_city.group(1)
            city = city[:zip_in_city.start()].strip()

    addr_text = ", ".join(rem).strip()

    # 3. Name vs Street Address Extraction
    words = addr_text.split()
    if len(words) >= 4:
        w1_low = words[1].lower()
        w2_low = words[2].lower()

        if re.search(r'\d', w1_low):
            contact_name = words[0]
            street_address = " ".join(words[1:])
        elif (re.search(r'\d', w2_low) or
              any(w2_low.endswith(sfx) for sfx in ADDR_SUFFIXES) or
              w2_low in ADDR_KEYWORDS):
            contact_name = " ".join(words[:2])
            street_address = " ".join(words[2:])
        else:
            w3_low = words[3].lower() if len(words) > 3 else ""
            if (w3_low and (re.search(r'\d', w3_low) or
                any(w3_low.endswith(sfx) for sfx in ADDR_SUFFIXES) or
                w3_low in ADDR_KEYWORDS)):
                contact_name = " ".join(words[:3])
                street_address = " ".join(words[3:])
            else:
                contact_name = " ".join(words[:2])
                street_address = " ".join(words[2:])
    elif len(words) == 3:
        if re.search(r'\d', words[1]):
            contact_name = words[0]
            street_address = " ".join(words[1:])
        elif re.search(r'\d', words[2]) or any(words[2].lower().endswith(sfx) for sfx in ADDR_SUFFIXES):
            contact_name = " ".join(words[:2])
            street_address = words[2]
        else:
            contact_name = " ".join(words[:2])
            street_address = words[2]
    elif len(words) == 2:
        contact_name = words[0]
        street_address = words[1]
    else:
        street_address = addr_text

    return {
        "handle": handle,
        "contact_name": contact_name,
        "street_address": street_address,
        "city": city,
        "postal_code": postal_code,
        "country": country.upper() if len(country) <= 3 else country.title(),
        "raw_disclosure": raw,
        "imported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }


def parse_vero_text(raw_text: str) -> List[Dict[str, str]]:
    """Parse multi-line string containing VeRO / eBay seller disclosures."""
    results = []
    seen_handles = set()
    if not raw_text:
        return results

    for line in raw_text.splitlines():
        line = line.strip()
        if not line or line.startswith("#") or line.startswith("//"):
            continue
        rec = parse_vero_line(line)
        if rec and rec["handle"]:
            h_key = rec["handle"].lower()
            if h_key not in seen_handles:
                seen_handles.add(h_key)
                results.append(rec)
    return results


def parse_vero_pdf(file_path: str) -> List[Dict[str, str]]:
    """Extract text from a VeRO PDF file and parse structured disclosure records."""
    if not PYPDF_AVAILABLE:
        raise RuntimeError("pypdf is required to read PDF files. Please install pypdf via 'pip install pypdf'.")

    if not os.path.exists(file_path):
        raise FileNotFoundError(f"PDF file not found: {file_path}")

    reader = PdfReader(file_path)
    full_text = []
    for page in reader.pages:
        txt = page.extract_text() or ""
        full_text.append(txt)

    combined_text = "\n".join(full_text)
    return parse_vero_text(combined_text)


def push_to_enforcement_registry(records: List[Dict[str, str]], data_store: Any) -> int:
    """
    Ingest parsed VeRO seller disclosures into Apollo's permanent Enforcement Registry.
    Updates seller metadata, country of origin, contact names, and structured address.
    """
    if not records or not data_store:
        return 0

    reg = data_store.get_enforcement_registry()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    updated_count = 0

    for rec in records:
        h = rec.get("handle")
        if not h:
            continue

        entry = reg.setdefault(h, {
            "seller_handle": h,
            "first_detected": now_str,
            "last_scanned": now_str,
            "total_infringements": 0,
            "total_clean_items": 0,
            "total_value": 0.0,
            "total_listings": 0,
            "items": [],
            "brands_targeted": [],
            "product_types": [],
            "locations": [],
            "country": "Unresolved",
            "threat_badge": "🛡 VeRO Disclosed Origin",
            "risk_tier": "Low",
            "historical_scan_count": 0,
            "strike_history": []
        })

        # Inject / Enrich disclosure intelligence
        if rec.get("contact_name"):
            entry["contact_name"] = rec["contact_name"]
        if rec.get("street_address"):
            entry["street_address"] = rec["street_address"]
        if rec.get("city"):
            entry["city"] = rec["city"]
        if rec.get("postal_code"):
            entry["postal_code"] = rec["postal_code"]
        if rec.get("country") and rec["country"] not in ("", "Unresolved"):
            entry["country"] = rec["country"]
        if rec.get("raw_disclosure"):
            entry["raw_disclosure"] = rec["raw_disclosure"]
        entry["disclosure_source"] = "VeRO / eBay Disclosure"
        entry["last_disclosure_import"] = now_str

        # Update threat badge
        if "CN" in str(entry.get("country", "")).upper() or "CHINA" in str(entry.get("country", "")).upper():
            entry["threat_badge"] = "🇨🇳 Verified Direct China Hub (VeRO)"
        elif not entry.get("threat_badge"):
            entry["threat_badge"] = "🛡 VeRO Disclosed Origin"

        updated_count += 1

    data_store.save_enforcement_registry(reg)
    logger.info(f"Successfully saved {updated_count} VeRO disclosures to Enforcement Registry.")
    return updated_count


def export_to_excel(records: List[Dict[str, str]], filepath: str):
    """
    Export parsed VeRO disclosures to a beautifully formatted Genesis-ready Excel spreadsheet.
    """
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl is required to export to Excel.")

    wb = openpyxl.Workbook()
    wb.properties.creator = "Jerry Seidenstucker"
    wb.properties.title = "Apollo VeRO Seller Disclosure Report"
    wb.remove(wb.active)

    ws = wb.create_sheet(title="VeRO Disclosures")

    # Header styling
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
    row_fill_a = PatternFill("solid", fgColor="FFFFFF")
    row_fill_b = PatternFill("solid", fgColor="F8FAFC")
    border_side = Side(style="thin", color="E2E8F0")
    thin_border = Border(bottom=border_side, top=border_side, left=border_side, right=border_side)

    headers = [
        "Seller Handle", "Contact / Legal Name", "Street Address", "City",
        "Postal / Zip Code", "Country", "Raw Disclosure Line", "Import Date", "Source"
    ]

    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 24

    for r_idx, rec in enumerate(records, 2):
        row_fill = row_fill_a if r_idx % 2 == 0 else row_fill_b
        row_data = [
            rec.get("handle", ""),
            rec.get("contact_name", ""),
            rec.get("street_address", ""),
            rec.get("city", ""),
            rec.get("postal_code", ""),
            rec.get("country", ""),
            rec.get("raw_disclosure", ""),
            rec.get("imported_at", ""),
            "VeRO / eBay Disclosure"
        ]
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="Segoe UI", size=9)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[r_idx].height = 20

    col_widths = [20, 22, 40, 18, 16, 12, 60, 20, 22]
    for idx, width in enumerate(col_widths, 1):
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "A2"
    wb.save(filepath)
    logger.info(f"VeRO Excel report exported: {filepath}")


def export_to_csv(records: List[Dict[str, str]], filepath: str):
    """Export parsed VeRO disclosures to CSV."""
    fieldnames = [
        "handle", "contact_name", "street_address", "city",
        "postal_code", "country", "raw_disclosure", "imported_at"
    ]
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in records:
            writer.writerow({k: r.get(k, "") for k in fieldnames})
    logger.info(f"VeRO CSV report exported: {filepath}")
